from dataclasses import dataclass
from decimal import Decimal
from decimal import InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook
import xlrd


ADD_HEADERS = ("商品名稱", "商品分類", "成本", "價錢", "庫存")
UPDATE_HEADERS = (
    "商品ID",
    "商品名稱",
    "商品分類",
    "成本",
    "價錢",
    "庫存",
)
PRICE_HEADERS = {"價錢", "價格"}


class ProductExcelError(ValueError):
    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.public_message = message


@dataclass(frozen=True, slots=True)
class ProductExcelRow:
    operation: str
    row_number: int
    product_id: str | None
    name: str
    label_names: str
    cost: Decimal
    price: Decimal
    stock: int


class ProductExcelService(object):
    ADD_DATA_ERROR = "新增商品資料錯誤，請檢查Excel內容"
    UPDATE_DATA_ERROR = "修改商品資料錯誤，請檢查Excel內容"
    ADD_HEADERS_ERROR = "新增商品欄位缺失，請使用正確的Excel模板"
    UPDATE_HEADERS_ERROR = "修改商品欄位缺失，請使用正確的Excel模板"
    INVALID_FILE_ERROR = "檔案格式不正確，請上傳Excel檔案"

    @classmethod
    def parse(
        cls,
        filename: str,
        content: bytes,
    ) -> list[ProductExcelRow]:
        suffix = Path(filename).suffix.casefold()
        if suffix not in {".xlsx", ".xls"} or not content:
            raise ProductExcelError(cls.INVALID_FILE_ERROR)
        try:
            if suffix == ".xlsx":
                sheets = cls.__read_xlsx(content)
            else:
                sheets = cls.__read_xls(content)
        except (
            OSError,
            ValueError,
            BadZipFile,
            InvalidFileException,
            KeyError,
            xlrd.XLRDError,
        ) as error:
            raise ProductExcelError(cls.INVALID_FILE_ERROR) from error

        rows: list[ProductExcelRow] = []
        for sheet_name, values in sheets.items():
            if sheet_name == "Add Product":
                rows.extend(cls.__parse_sheet(values, "add"))
            elif sheet_name == "Update Product":
                rows.extend(cls.__parse_sheet(values, "update"))
        return rows

    @staticmethod
    def create_workbook(
        products: list[dict[str, Any]],
    ) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Update Product"
        sheet.append(UPDATE_HEADERS)
        for product in products:
            sheet.append(
                [
                    product["id"],
                    product["name"],
                    product["label_names"],
                    product["cost"],
                    product["price"],
                    product["stock"],
                ]
            )
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def create_template() -> bytes:
        workbook = Workbook()
        add_sheet = workbook.active
        add_sheet.title = "Add Product"
        add_sheet.append(ADD_HEADERS)
        update_sheet = workbook.create_sheet("Update Product")
        update_sheet.append(UPDATE_HEADERS)
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @classmethod
    def __read_xlsx(cls, content: bytes) -> dict[str, list[list[Any]]]:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            return {
                sheet.title: [
                    [cell.value for cell in row]
                    for row in sheet.iter_rows()
                ]
                for sheet in workbook.worksheets
            }
        finally:
            workbook.close()

    @staticmethod
    def __read_xls(content: bytes) -> dict[str, list[list[Any]]]:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            return {
                sheet.name: [
                    sheet.row_values(row_index)
                    for row_index in range(sheet.nrows)
                ]
                for sheet in workbook.sheets()
            }
        finally:
            workbook.release_resources()

    @classmethod
    def __parse_sheet(
        cls,
        values: list[list[Any]],
        operation: str,
    ) -> list[ProductExcelRow]:
        expected_headers = ADD_HEADERS if operation == "add" else UPDATE_HEADERS
        header_error = (
            cls.ADD_HEADERS_ERROR
            if operation == "add"
            else cls.UPDATE_HEADERS_ERROR
        )
        data_error = (
            cls.ADD_DATA_ERROR
            if operation == "add"
            else cls.UPDATE_DATA_ERROR
        )
        if not values:
            raise ProductExcelError(header_error)

        headers = [cls.__normalize_header(value) for value in values[0]]
        header_positions: dict[str, int] = {}
        for index, header in enumerate(headers):
            normalized = "價錢" if header in PRICE_HEADERS else header
            if normalized in expected_headers:
                if normalized in header_positions:
                    raise ProductExcelError(header_error)
                header_positions[normalized] = index
        if any(header not in header_positions for header in expected_headers):
            raise ProductExcelError(header_error)

        products: list[ProductExcelRow] = []
        for row_number, raw_row in enumerate(values[1:], start=2):
            if not any(value not in (None, "") for value in raw_row):
                continue
            try:
                values_by_header = {
                    header: cls.__cell_value(raw_row, index)
                    for header, index in header_positions.items()
                }
                product_id = (
                    cls.__product_id(values_by_header["商品ID"])
                    if operation == "update"
                    else None
                )
                name = cls.__text(values_by_header["商品名稱"])
                label_names = cls.__text(values_by_header["商品分類"])
                cost = cls.__decimal(values_by_header["成本"])
                price = cls.__decimal(values_by_header["價錢"])
                stock_value = values_by_header["庫存"]
                stock = 0 if stock_value in (None, "") else cls.__integer(stock_value)
                if (
                    not name
                    or len(name) > 255
                    or not label_names
                    or any(
                        len(label.strip()) > 255
                        for label in label_names.split(",")
                    )
                    or cost < 0
                    or price < 0
                    or stock < 0
                    or stock > 2_147_483_647
                ):
                    raise ValueError
                products.append(
                    ProductExcelRow(
                        operation=operation,
                        row_number=row_number,
                        product_id=product_id,
                        name=name,
                        label_names=label_names,
                        cost=cost,
                        price=price,
                        stock=stock,
                    )
                )
            except (InvalidOperation, TypeError, ValueError, KeyError) as error:
                raise ProductExcelError(data_error) from error
        return products

    @staticmethod
    def __normalize_header(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def __cell_value(row: list[Any], index: int) -> Any:
        return row[index] if index < len(row) else None

    @staticmethod
    def __text(value: Any) -> str:
        if value is None:
            raise ValueError
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @classmethod
    def __product_id(cls, value: Any) -> str:
        product_id = cls.__text(value)
        if product_id.endswith(".0"):
            product_id = product_id[:-2]
        if len(product_id) != 13 or not product_id.isdigit():
            raise ValueError
        return product_id

    @classmethod
    def __decimal(cls, value: Any) -> Decimal:
        if value is None or isinstance(value, bool):
            raise ValueError
        amount = Decimal(str(value))
        if (
            not amount.is_finite()
            or amount.as_tuple().exponent < -2
            or abs(amount) >= Decimal("100000000")
        ):
            raise ValueError
        return amount

    @classmethod
    def __integer(cls, value: Any) -> int:
        if isinstance(value, bool):
            raise ValueError
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError
        return int(number)


