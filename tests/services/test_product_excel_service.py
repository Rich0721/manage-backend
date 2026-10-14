from io import BytesIO
from pathlib import Path
import struct
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import pytest
import xlrd

from src.services.product_excel_service import ProductExcelError
from src.services.product_excel_service import ProductExcelService


def make_workbook_bytes() -> bytes:
    workbook = Workbook()
    add_sheet = workbook.active
    add_sheet.title = "Add Product"
    add_sheet.append(["商品名稱", "商品分類", "成本", "價格", "庫存"])
    add_sheet.append(["New", "Label A", 10, 20, None])
    update_sheet = workbook.create_sheet("Update Product")
    update_sheet.append(
        ["商品ID", "商品名稱", "商品分類", "成本", "價錢", "庫存"]
    )
    update_sheet.append(
        ["1790705105001", "Existing", "Label B", 30, 40, 5]
    )
    workbook.create_sheet("Ignored")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_biff_workbook_bytes(
    *,
    formula_in_name: bool = False,
    error_in_cost: bool = False,
    biff_version: int = 0x0600,
) -> bytes:
    def record(record_id: int, data: bytes) -> bytes:
        return struct.pack("<HH", record_id, len(data)) + data

    def bof(stream_type: int) -> bytes:
        data = struct.pack(
            "<HHHHII",
            biff_version,
            stream_type,
            0x0DBB,
            0x07CC,
            0x00000041,
            0x00000006,
        )
        return record(0x0809, data)

    def label(row: int, column: int, value: str) -> bytes:
        if biff_version < 0x0600:
            encoded_value = value.encode("cp950")
            data = (
                struct.pack("<HHHH", row, column, 0, len(encoded_value))
                + encoded_value
            )
        else:
            data = (
                struct.pack("<HHHHB", row, column, 0, len(value), 1)
                + value.encode("utf-16le")
            )
        return record(0x0204, data)

    codepage = 950 if biff_version < 0x0600 else 1200
    global_prefix = bof(0x0005) + record(
        0x0042,
        struct.pack("<H", codepage),
    )
    sheet_name = b"Add Product"

    def boundsheet_data(offset: int) -> bytes:
        data = struct.pack("<I", offset) + b"\0\0" + bytes([len(sheet_name)])
        if biff_version >= 0x0600:
            data += b"\0"
        return data + sheet_name

    boundsheet_record = record(0x0085, boundsheet_data(0))
    sheet_offset = len(global_prefix) + len(boundsheet_record) + 4
    workbook_globals = (
        global_prefix
        + record(0x0085, boundsheet_data(sheet_offset))
        + record(0x000A, b"")
    )

    sheet_records = [
        bof(0x0010),
        label(0, 0, "商品名稱"),
        label(0, 1, "商品分類"),
        label(0, 2, "成本"),
        label(0, 3, "價錢"),
        label(0, 4, "庫存"),
    ]
    if formula_in_name:
        formula_data = (
            struct.pack("<HHH", 1, 0, 0)
            + struct.pack("<d", 2)
            + b"\0" * 6
            + struct.pack("<H", 0)
        )
        sheet_records.append(record(0x0006, formula_data))
    else:
        sheet_records.append(label(1, 0, "Legacy product"))
    sheet_records.append(label(1, 1, "Legacy label"))
    if error_in_cost:
        sheet_records.append(
            record(0x0205, struct.pack("<HHHBB", 1, 2, 0, 7, 1))
        )
    else:
        sheet_records.append(record(0x0203, struct.pack("<HHHd", 1, 2, 0, 10)))
    sheet_records.extend(
        [
            record(0x0203, struct.pack("<HHHd", 1, 3, 0, 20)),
            record(0x0203, struct.pack("<HHHd", 1, 4, 0, 3)),
            record(0x000A, b""),
        ]
    )
    return workbook_globals + b"".join(sheet_records)


def test_parse_reads_a_real_raw_biff_xls_workbook() -> None:
    rows = ProductExcelService.parse(
        "products.xls",
        make_biff_workbook_bytes(),
    )

    assert len(rows) == 1
    assert rows[0].name == "Legacy product"
    assert rows[0].label_names == "Legacy label"
    assert rows[0].cost == 10
    assert rows[0].price == 20
    assert rows[0].stock == 3


def test_parse_reads_a_real_biff5_xls_workbook() -> None:
    rows = ProductExcelService.parse(
        "products.xls",
        make_biff_workbook_bytes(biff_version=0x0500),
    )

    assert len(rows) == 1
    assert rows[0].name == "Legacy product"
    assert rows[0].label_names == "Legacy label"
    assert rows[0].cost == 10
    assert rows[0].price == 20
    assert rows[0].stock == 3


def test_parse_rejects_formula_in_a_real_raw_biff_xls_workbook() -> None:
    with pytest.raises(ProductExcelError) as error:
        ProductExcelService.parse(
            "products.xls",
            make_biff_workbook_bytes(formula_in_name=True),
        )

    assert error.value.public_message == ProductExcelService.ADD_DATA_ERROR


def test_parse_rejects_formula_in_a_real_biff5_xls_workbook() -> None:
    with pytest.raises(ProductExcelError) as error:
        ProductExcelService.parse(
            "products.xls",
            make_biff_workbook_bytes(
                formula_in_name=True,
                biff_version=0x0500,
            ),
        )

    assert error.value.public_message == ProductExcelService.ADD_DATA_ERROR


def test_parse_rejects_error_cell_in_a_real_raw_biff_xls_workbook() -> None:
    with pytest.raises(ProductExcelError) as error:
        ProductExcelService.parse(
            "products.xls",
            make_biff_workbook_bytes(error_in_cost=True),
        )

    assert error.value.public_message == ProductExcelService.ADD_DATA_ERROR


def test_parse_rejects_error_cell_in_a_real_biff5_xls_workbook() -> None:
    with pytest.raises(ProductExcelError) as error:
        ProductExcelService.parse(
            "products.xls",
            make_biff_workbook_bytes(
                error_in_cost=True,
                biff_version=0x0500,
            ),
        )

    assert error.value.public_message == ProductExcelService.ADD_DATA_ERROR


def test_xls_formula_scanner_reads_legacy_book_stream() -> None:
    biff_stream = make_biff_workbook_bytes(
        formula_in_name=True,
        biff_version=0x0500,
    )

    class FakeCompoundDocument(object):
        def __init__(self, content: bytes, logfile: object) -> None:
            return None

        @staticmethod
        def get_named_stream(name: str) -> bytes | None:
            if name == "Book":
                return biff_stream
            return None

    content = xlrd.compdoc.SIGNATURE + b"unused"
    with patch(
        "src.services.product_excel_service.xlrd.compdoc.CompDoc",
        FakeCompoundDocument,
    ):
        formulas = ProductExcelService._ProductExcelService__xls_formula_cells(
            content
        )

    assert formulas == {"Add Product": {(1, 0)}}


def test_parse_xlsx_sheets_and_price_alias() -> None:
    rows = ProductExcelService.parse("products.xlsx", make_workbook_bytes())

    assert len(rows) == 2
    assert rows[0].operation == "add"
    assert rows[0].stock == 0
    assert rows[0].price == 20
    assert rows[1].operation == "update"
    assert rows[1].product_id == "1790705105001"


def test_parse_rejects_missing_add_headers() -> None:
    workbook = Workbook()
    workbook.active.title = "Add Product"
    workbook.active.append(["商品名稱"])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ProductExcelError) as error:
        ProductExcelService.parse("products.xlsx", output.getvalue())

    assert error.value.public_message == ProductExcelService.ADD_HEADERS_ERROR


def test_parse_rejects_invalid_update_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Update Product"
    sheet.append(["商品ID", "商品名稱", "商品分類", "成本", "價錢", "庫存"])
    sheet.append(["1790705105001", "Bad", "Label", 10, 20, -1])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ProductExcelError) as error:
        ProductExcelService.parse("products.xlsx", output.getvalue())

    assert error.value.public_message == ProductExcelService.UPDATE_DATA_ERROR


@pytest.mark.parametrize("invalid_value", ["=1+1", "#VALUE!"])
def test_parse_rejects_formula_and_error_cells(invalid_value: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Add Product"
    sheet.append(["商品名稱", "商品分類", "成本", "價錢", "庫存"])
    sheet.append([invalid_value, "Label", 10, 20, 0])
    if invalid_value.startswith("#"):
        sheet["A2"].data_type = "e"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ProductExcelError) as error:
        ProductExcelService.parse("products.xlsx", output.getvalue())

    assert error.value.public_message == ProductExcelService.ADD_DATA_ERROR


def test_parse_reads_legacy_xls_sheets() -> None:
    class FakeSheet(object):
        name = "Add Product"
        nrows = 2

        class Cell(object):
            def __init__(self, value: object) -> None:
                self.value = value
                self.ctype = 1

        @classmethod
        def row(cls, row_index: int) -> list[object]:
            return [cls.Cell(value) for value in cls.row_values(row_index)]

        @staticmethod
        def row_values(row_index: int) -> list[object]:
            if row_index == 0:
                return ["商品名稱", "商品分類", "成本", "價錢", "庫存"]
            return ["Legacy", "Label", 10, 20, 2]

    class FakeWorkbook(object):
        @staticmethod
        def sheets() -> list[FakeSheet]:
            return [FakeSheet()]

        @staticmethod
        def release_resources() -> None:
            return None

    class FakeCompoundDocument(object):
        @staticmethod
        def get_named_stream(name: str) -> bytes:
            return b""

    with (
        patch(
            "src.services.product_excel_service.xlrd.open_workbook",
            return_value=FakeWorkbook(),
        ),
        patch(
            "src.services.product_excel_service.xlrd.compdoc.CompDoc",
            return_value=FakeCompoundDocument(),
        ),
    ):
        rows = ProductExcelService.parse(
            "products.xls",
            xlrd.compdoc.SIGNATURE + b"xls-bytes",
        )

    assert rows[0].name == "Legacy"
    assert rows[0].stock == 2


def test_parse_rejects_formula_in_legacy_xls() -> None:
    class FakeCell(object):
        def __init__(self, value: object) -> None:
            self.value = value
            self.ctype = 1

    class FakeSheet(object):
        name = "Add Product"
        nrows = 2

        @staticmethod
        def row(row_index: int) -> list[FakeCell]:
            values = (
                ["商品名稱", "商品分類", "成本", "價錢", "庫存"]
                if row_index == 0
                else ["Calculated formula", "Label", 10, 20, 0]
            )
            return [FakeCell(value) for value in values]

    class FakeWorkbook(object):
        @staticmethod
        def sheets() -> list[FakeSheet]:
            return [FakeSheet()]

        @staticmethod
        def release_resources() -> None:
            return None

    def record(record_id: int, data: bytes) -> bytes:
        return struct.pack("<HH", record_id, len(data)) + data

    name = b"Add Product"
    boundsheet = record(0x0085, b"\0\0\0\0\0\0" + bytes([len(name), 0]) + name)
    sheet_offset = len(boundsheet)
    boundsheet_data = (
        struct.pack("<I", sheet_offset)
        + b"\0\0"
        + bytes([len(name), 0])
        + name
    )
    formula_stream = (
        record(0x0085, boundsheet_data)
        + record(0x0809, b"")
        + record(0x0006, struct.pack("<HH", 1, 0) + bytes(20))
        + record(0x000A, b"")
    )

    class FakeCompoundDocument(object):
        def __init__(self, content: bytes, logfile: object) -> None:
            return None

        @staticmethod
        def get_named_stream(name: str) -> bytes:
            return formula_stream

    with (
        patch(
            "src.services.product_excel_service.xlrd.open_workbook",
            return_value=FakeWorkbook(),
        ),
        patch(
            "src.services.product_excel_service.xlrd.compdoc.CompDoc",
            FakeCompoundDocument,
        ),
        pytest.raises(ProductExcelError) as error,
    ):
        ProductExcelService.parse(
            "products.xls",
            xlrd.compdoc.SIGNATURE + b"xls-bytes",
        )

    assert error.value.public_message == ProductExcelService.ADD_DATA_ERROR


def test_parse_rejects_error_cell_in_legacy_xls() -> None:
    class FakeCell(object):
        def __init__(self, value: object, cell_type: int = 1) -> None:
            self.value = value
            self.ctype = cell_type

    class FakeSheet(object):
        name = "Add Product"
        nrows = 2

        @staticmethod
        def row(row_index: int) -> list[FakeCell]:
            values = (
                ["商品名稱", "商品分類", "成本", "價錢", "庫存"]
                if row_index == 0
                else ["Bad", "Label", "#VALUE!", 20, 0]
            )
            return [
                FakeCell(value, 5 if value == "#VALUE!" else 1)
                for value in values
            ]

    class FakeWorkbook(object):
        @staticmethod
        def sheets() -> list[FakeSheet]:
            return [FakeSheet()]

        @staticmethod
        def release_resources() -> None:
            return None

    class FakeCompoundDocument(object):
        @staticmethod
        def get_named_stream(name: str) -> bytes:
            return b""

    with (
        patch(
            "src.services.product_excel_service.xlrd.open_workbook",
            return_value=FakeWorkbook(),
        ),
        patch(
            "src.services.product_excel_service.xlrd.compdoc.CompDoc",
            return_value=FakeCompoundDocument(),
        ),
        pytest.raises(ProductExcelError) as error,
    ):
        ProductExcelService.parse(
            "products.xls",
            xlrd.compdoc.SIGNATURE + b"xls-bytes",
        )

    assert error.value.public_message == ProductExcelService.ADD_DATA_ERROR


def test_template_and_export_have_required_sheets_and_headers() -> None:
    template = load_workbook(BytesIO(ProductExcelService.create_template()))
    exported = load_workbook(
        BytesIO(ProductExcelService.create_workbook([]))
    )

    assert template.sheetnames == ["Add Product", "Update Product"]
    assert list(template["Add Product"].values)[0] == (
        "商品名稱",
        "商品分類",
        "成本",
        "價錢",
        "庫存",
    )
    assert exported.sheetnames == ["Update Product"]
    assert list(exported["Update Product"].values)[0] == (
        "商品ID",
        "商品名稱",
        "商品分類",
        "成本",
        "價錢",
        "庫存",
    )


def test_static_template_is_readable_and_contains_both_operations() -> None:
    template_path = Path("files/Batch Add Products.xlsx")
    workbook = load_workbook(template_path, read_only=True, data_only=True)

    assert workbook.sheetnames == ["Add Product", "Update Product"]
    assert next(workbook["Add Product"].values) == (
        "商品名稱",
        "商品分類",
        "成本",
        "價錢",
        "庫存",
    )
    assert next(workbook["Update Product"].values) == (
        "商品ID",
        "商品名稱",
        "商品分類",
        "成本",
        "價錢",
        "庫存",
    )
    workbook.close()
