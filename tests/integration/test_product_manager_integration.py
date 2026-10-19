import asyncio
import json
import os
from io import BytesIO
from contextlib import asynccontextmanager
from dataclasses import replace
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from psycopg import AsyncConnection
from redis.asyncio import Redis

from src.constants.user import UserRole
from src.models.po.label import Label
from src.models.po.product import Product
from src.models.po.user import UserPO
from src.models.schemas.authorization import AuthorizationObject
from src.models.schemas.product import ProductCreateInfo
from src.models.schemas.product import ProductDownloadInfo
from src.repositories.product_cache_repository import ProductCacheRepository
from src.repositories.product_cache_repository import ProductCacheRepositoryError
from src.repositories.label_cache_repository import LabelCacheRepository
from src.repositories.label_cache_repository import LabelCacheRepositoryError
from src.repositories.label_repository import LabelRepository
from src.repositories.product_repository import ProductRepository
from src.repositories.product_repository import ProductRepositoryError
from src.services.authorization_service import AuthorizationContext
from src.services.errors import ServiceUnavailableError
from src.services.product_errors import ProductCacheUnavailableError
from src.services.product_errors import ProductNotFoundError
from src.services.product_service import ProductService


class DirectConnectionPool(object):
    """Expose a test connection through the repository pool interface."""

    def __init__(self, connection: AsyncConnection) -> None:
        self.connection_value = connection

    @asynccontextmanager
    async def connection(self):
        yield self.connection_value


class IntegrationAuthorization(object):
    async def validate_session(
        self,
        auth: AuthorizationObject,
    ) -> AuthorizationContext:
        return AuthorizationContext(auth.uid or "integration-user", None, True)

    async def refresh_session(self, context: AuthorizationContext) -> None:
        return None


class IntegrationUserRepository(object):
    async def get_by_uid(self, uid: str) -> UserPO:
        now = datetime(2026, 1, 1)
        return UserPO(
            uid=uid,
            email="integration@example.com",
            user_name="Integration",
            password="unused",
            permission=UserRole.ADMIN,
            created_at=now,
            updated_at=now,
        )


def create_add_workbook(name: str, label: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Add Product"
    sheet.append(["商品名稱", "商品分類", "成本", "價錢", "庫存"])
    sheet.append([name, label, 10, 20, 3])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_mixed_workbook() -> bytes:
    workbook = Workbook()
    add_sheet = workbook.active
    add_sheet.title = "Add Product"
    add_sheet.append(["商品名稱", "商品分類", "成本", "價錢", "庫存"])
    add_sheet.append(["Must rollback", "New label", 10, 20, 3])
    update_sheet = workbook.create_sheet("Update Product")
    update_sheet.append(["商品ID", "商品名稱", "商品分類", "成本", "價錢", "庫存"])
    update_sheet.append(["1790705105999", "Missing", "New label", 10, 20, 3])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_integration_product_service(
    connection: AsyncConnection,
    redis: Redis,
    prefix: str,
) -> tuple[ProductService, ProductRepository, LabelRepository,
           ProductCacheRepository, LabelCacheRepository]:
    pool = DirectConnectionPool(connection)
    products = ProductRepository(pool)  # type: ignore[arg-type]
    labels = LabelRepository(pool)  # type: ignore[arg-type]
    product_cache = ProductCacheRepository(
        redis,
        f"test:{prefix}:products:info",
    )
    label_cache = LabelCacheRepository(
        redis,
        f"test:{prefix}:product:labels",
    )
    service = ProductService(
        object(),  # type: ignore[arg-type]
        products,
        product_cache,
        labels,
        label_cache,
        IntegrationUserRepository(),  # type: ignore[arg-type]
        IntegrationAuthorization(),  # type: ignore[arg-type]
    )
    return service, products, labels, product_cache, label_cache


async def create_temp_product_tables(
    connection: AsyncConnection,
) -> None:
    labels_ddl = LABEL_DDL_PATH.read_text(encoding="utf-8").replace(
        "CREATE TABLE tb_labels",
        "CREATE TEMP TABLE tb_labels",
        1,
    )
    products_ddl = PRODUCT_DDL_PATH.read_text(encoding="utf-8").replace(
        "CREATE TABLE tb_products",
        "CREATE TEMP TABLE tb_products",
        1,
    )
    await connection.execute(labels_ddl)
    await connection.execute(products_ddl)
    await connection.commit()


DATABASE_URL = os.getenv("INTEGRATION_DATABASE_URL")
REDIS_URL = os.getenv("INTEGRATION_REDIS_URL")
LABEL_DDL_PATH = Path("database/DDL/tables/TB_LABELS.sql")
PRODUCT_DDL_PATH = Path("database/DDL/tables/TB_PRODUCTS.sql")
PRODUCT_STOCK_MIGRATION_PATH = Path(
    "database/DDL/tables/ALTER_TB_PRODUCTS_ADD_STOCK.sql"
)


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None,
    reason="INTEGRATION_DATABASE_URL is not configured",
)
async def test_product_and_label_ddl_support_soft_deleted_products() -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    try:
        labels_ddl = LABEL_DDL_PATH.read_text(encoding="utf-8").replace(
            "CREATE TABLE tb_labels",
            "CREATE TEMP TABLE tb_labels",
            1,
        )
        products_ddl = PRODUCT_DDL_PATH.read_text(encoding="utf-8").replace(
            "CREATE TABLE tb_products",
            "CREATE TEMP TABLE tb_products",
            1,
        )
        await connection.execute(labels_ddl)
        await connection.execute(products_ddl)
        now = datetime(2026, 1, 1)
        await connection.execute(
            """
            INSERT INTO tb_labels (
                id, name, created_uid, created_at, updated_uid, updated_at
            ) VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (1, "Label", "user", now, "user", now),
        )
        await connection.execute(
            """
            INSERT INTO tb_products (
                id, name, label_ids, cost, price, delete_flag,
                created_uid, created_at, updated_uid, updated_at
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                "1790705105001",
                "Deleted",
                "1",
                Decimal("100.00"),
                Decimal("150.00"),
                True,
                "user",
                now,
                "user",
                now,
            ),
        )
        cursor = await connection.execute(
            "SELECT delete_flag FROM tb_products WHERE id = %s",
            ("1790705105001",),
        )
        row = await cursor.fetchone()
        assert row == (True,)
    finally:
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None,
    reason="INTEGRATION_DATABASE_URL is not configured",
)
async def test_stock_migration_updates_existing_and_new_rows() -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    try:
        await connection.execute(
            "CREATE TEMP TABLE tb_products (id VARCHAR(13) PRIMARY KEY)"
        )
        await connection.execute(
            "INSERT INTO tb_products (id) VALUES (%s)",
            ("1790705105001",),
        )
        await connection.commit()

        migration = PRODUCT_STOCK_MIGRATION_PATH.read_text(encoding="utf-8")
        await connection.execute(migration)
        await connection.execute(migration)
        await connection.execute(
            "INSERT INTO tb_products (id) VALUES (%s)",
            ("1790705105002",),
        )
        cursor = await connection.execute(
            "SELECT stock FROM tb_products ORDER BY id"
        )

        assert await cursor.fetchall() == [(0,), (0,)]
    finally:
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None,
    reason="INTEGRATION_DATABASE_URL is not configured",
)
async def test_product_repository_transaction_rolls_back_on_failure() -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    try:
        products_ddl = PRODUCT_DDL_PATH.read_text(encoding="utf-8").replace(
            "CREATE TABLE tb_products",
            "CREATE TEMP TABLE tb_products",
            1,
        )
        await connection.execute(products_ddl)
        repository = ProductRepository(DirectConnectionPool(connection))  # type: ignore[arg-type]
        product = Product(
            id="1790705105001",
            name="Rollback",
            label_ids="1",
            cost=Decimal("100.00"),
            price=Decimal("150.00"),
            delete_flag=False,
            created_uid="user",
            created_at=datetime(2026, 1, 1),
            updated_uid="user",
            updated_at=datetime(2026, 1, 1),
        )

        with pytest.raises(RuntimeError):
            async with repository.transaction() as transaction:
                await repository.insert(product, transaction)
                raise RuntimeError("force rollback")

        assert await repository.get_by_id(product.id) is None
    finally:
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    REDIS_URL is None,
    reason="INTEGRATION_REDIS_URL is not configured",
)
async def test_products_cache_preserves_complete_snapshot() -> None:
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    cache = ProductCacheRepository(client, f"test:{prefix}:products:info")
    product = Product(
        id="1790705105001",
        name=f"Deleted-{prefix}",
        label_ids="1",
        cost=Decimal("100.00"),
        price=Decimal("150.00"),
        delete_flag=True,
        created_uid="user",
        created_at=datetime(2026, 1, 1),
        updated_uid="user",
        updated_at=datetime(2026, 1, 2),
    )
    try:
        async with cache.products_lock() as lease:
            assert await client.ttl(f"test:{prefix}:products:info:lock") > 0
            await cache.replace_products([product], {1: "Label"}, lease)
        cached_products = await cache.get_products()
        assert cached_products is not None
        assert cached_products[0].label_names == ("Label",)
    finally:
        await client.delete(
            f"test:{prefix}:products:info",
            f"test:{prefix}:products:info:version",
        )
        await client.aclose()


@pytest.mark.asyncio
@pytest.mark.skipif(
    REDIS_URL is None,
    reason="INTEGRATION_REDIS_URL is not configured",
)
async def test_label_cache_fences_old_writer_after_invalidation() -> None:
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    key = f"test:{prefix}:product:labels"
    cache = LabelCacheRepository(client, key)
    timestamp = datetime(2026, 1, 1)
    label = Label(1, "Label", "user", timestamp, "user", timestamp)
    try:
        async with cache.lock() as stale_lease:
            await cache.invalidate_after_commit()
            with pytest.raises(LabelCacheRepositoryError):
                await cache.replace([label], stale_lease)

        assert await cache.get() is None
    finally:
        await client.delete(key, f"{key}:version", f"{key}:lock")
        await client.aclose()


@pytest.mark.asyncio
@pytest.mark.skipif(
    REDIS_URL is None,
    reason="INTEGRATION_REDIS_URL is not configured",
)
async def test_products_cache_does_not_publish_after_losing_lease() -> None:
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    key = f"test:{prefix}:products:info"
    cache = ProductCacheRepository(client, key)
    product = Product(
        id="1790705105001",
        name="Old snapshot",
        label_ids="1",
        cost=Decimal("100.00"),
        price=Decimal("150.00"),
        delete_flag=False,
        created_uid="user",
        created_at=datetime(2026, 1, 1),
        updated_uid="user",
        updated_at=datetime(2026, 1, 2),
    )
    try:
        async with cache.products_lock() as lease:
            await client.set(f"{key}:lock", "new-owner")
            await client.set(key, "newer snapshot")
            with pytest.raises(ProductCacheRepositoryError):
                await cache.replace_products([product], {1: "Label"}, lease)

        assert await client.get(key) == "newer snapshot"
    finally:
        await client.delete(key, f"{key}:lock", f"{key}:version")
        await client.aclose()


@pytest.mark.asyncio
@pytest.mark.skipif(
    REDIS_URL is None,
    reason="INTEGRATION_REDIS_URL is not configured",
)
async def test_redis_fences_writer_and_cold_loader_after_commit() -> None:
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    key = f"test:{uuid4().hex}:products:info"
    cache = ProductCacheRepository(client, key)
    old_product = Product(
        id="1790705105001",
        name="Old snapshot",
        label_ids="1",
        cost=Decimal("100.00"),
        price=Decimal("150.00"),
        delete_flag=False,
        created_uid="user",
        created_at=datetime(2026, 1, 1),
        updated_uid="user",
        updated_at=datetime(2026, 1, 2),
    )
    committed_product = replace(old_product, name="Committed")
    try:
        async with cache.products_lock() as first_writer:
            await cache.invalidate_products(first_writer)
            await client.delete(f"{key}:lock")  # Simulate lease expiry.
            async with cache.products_lock() as second_writer:
                await cache.invalidate_products(second_writer)
                await cache.invalidate_after_commit()
                with pytest.raises(ProductCacheRepositoryError):
                    await cache.replace_products(
                        [old_product], {1: "Label"}, second_writer
                    )
            with pytest.raises(ProductCacheRepositoryError):
                await cache.replace_products(
                    [old_product], {1: "Label"}, first_writer
                )

        async with cache.products_lock() as cold_loader:
            assert await cache.get_products() is None
            await cache.replace_products(
                [committed_product], {1: "Label"}, cold_loader
            )
        products = await cache.get_products()
        assert products is not None
        assert products[0].name == "Committed"
    finally:
        await client.delete(key, f"{key}:lock", f"{key}:version")
        await client.aclose()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None or REDIS_URL is None,
    reason="Integration PostgreSQL and Redis URLs are not configured",
)
async def test_batch_products_commits_and_downloads_real_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    service, products, labels, product_cache, label_cache = (
        create_integration_product_service(connection, client, prefix)
    )
    try:
        await create_temp_product_tables(connection)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Add Product"
        sheet.append(["商品名稱", "商品分類", "成本", "價錢", "庫存"])
        sheet.append(["Batch one", "New label", 10, 20, 3])
        sheet.append(["Batch two", "New label", 11, 21, 4])
        output = BytesIO()
        workbook.save(output)

        result = await service.batch_products(
            AuthorizationObject(uid="integration-user"),
            "products.xlsx",
            output.getvalue(),
        )

        assert len(result.info.root) == 2
        product_rows = await products.list_all()
        label_rows = await labels.list_all()
        assert [product.name for product in product_rows] == [
            "Batch one",
            "Batch two",
        ]
        assert [product.stock for product in product_rows] == [3, 4]
        assert [label.name for label in label_rows] == ["New label"]
        assert await product_cache.get_products() is not None
        assert await label_cache.get() == {"new label": (1, "New label")}

        monkeypatch.setattr(
            "src.services.product_service.FILES_DIRECTORY",
            tmp_path,
        )
        first_id, second_id = [product.id for product in result.info.root]
        first_download, second_download = await asyncio.gather(
            service.download_product_data(
                AuthorizationObject(uid="integration-user"),
                ProductDownloadInfo(ids=first_id),
            ),
            service.download_product_data(
                AuthorizationObject(uid="integration-user"),
                ProductDownloadInfo(ids=second_id),
            ),
        )
        assert first_download.filename == second_download.filename
        saved_workbook = load_workbook(
            tmp_path / first_download.filename,
            read_only=True,
        )
        assert saved_workbook["Update Product"]["B2"].value in {
            "Batch one",
            "Batch two",
        }
        saved_workbook.close()
        for download, expected_name in (
            (first_download, "Batch one"),
            (second_download, "Batch two"),
        ):
            exported = load_workbook(BytesIO(download.content), read_only=True)
            assert exported["Update Product"]["B2"].value == expected_name
            exported.close()
    finally:
        await client.delete(
            f"test:{prefix}:products:info",
            f"test:{prefix}:products:info:version",
            f"test:{prefix}:products:info:lock",
            f"test:{prefix}:product:labels",
            f"test:{prefix}:product:labels:version",
            f"test:{prefix}:product:labels:lock",
        )
        await client.aclose()
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None or REDIS_URL is None,
    reason="Integration PostgreSQL and Redis URLs are not configured",
)
async def test_mixed_batch_rolls_back_database_and_keeps_caches() -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    service, products, labels, product_cache, label_cache = (
        create_integration_product_service(connection, client, prefix)
    )
    try:
        await create_temp_product_tables(connection)
        now = datetime(2026, 1, 1)
        async with products.transaction() as transaction:
            label = await labels.insert(
                Label(0, "Existing", "seed", now, "seed", now),
                transaction,
            )
            original = Product(
                id="1790705105001",
                name="Original",
                label_ids=str(label.id),
                cost=Decimal("10.00"),
                price=Decimal("20.00"),
                stock=8,
                delete_flag=False,
                created_uid="seed",
                created_at=now,
                updated_uid="seed",
                updated_at=now,
            )
            await products.insert(original, transaction)
        async with label_cache.lock() as label_lease:
            await label_cache.replace([label], label_lease)
        async with product_cache.products_lock() as product_lease:
            await product_cache.replace_products(
                [original], {label.id: label.name}, product_lease
            )

        with pytest.raises(ProductNotFoundError):
            await service.batch_products(
                AuthorizationObject(uid="integration-user"),
                "mixed.xlsx",
                create_mixed_workbook(),
            )

        assert [product.name for product in await products.list_all()] == [
            "Original",
        ]
        stored_labels = await labels.list_all()
        assert [label.name for label in stored_labels] == ["Existing"]
        assert (await product_cache.get_products())[0].stock == 8
        assert await label_cache.get() == {"existing": (label.id, label.name)}
    finally:
        await client.delete(
            f"test:{prefix}:products:info",
            f"test:{prefix}:products:info:version",
            f"test:{prefix}:products:info:lock",
            f"test:{prefix}:product:labels",
            f"test:{prefix}:product:labels:version",
            f"test:{prefix}:product:labels:lock",
        )
        await client.aclose()
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None or REDIS_URL is None,
    reason="Integration PostgreSQL and Redis URLs are not configured",
)
async def test_post_commit_label_cache_failure_invalidates_both_caches(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    service, products, labels, product_cache, label_cache = (
        create_integration_product_service(connection, client, prefix)
    )

    async def fail_replace(*args: object, **kwargs: object) -> None:
        raise LabelCacheRepositoryError("injected label cache write failure")

    monkeypatch.setattr(label_cache, "replace", fail_replace)
    try:
        await create_temp_product_tables(connection)
        with pytest.raises(ProductCacheUnavailableError):
            await service.batch_products(
                AuthorizationObject(uid="integration-user"),
                "products.xlsx",
                create_add_workbook("Committed product", "New label"),
            )

        assert len(await products.list_all()) == 1
        assert len(await labels.list_all()) == 1
        assert await product_cache.get_products() is None
        assert await label_cache.get() is None
    finally:
        await client.delete(
            f"test:{prefix}:products:info",
            f"test:{prefix}:products:info:version",
            f"test:{prefix}:products:info:lock",
            f"test:{prefix}:product:labels",
            f"test:{prefix}:product:labels:version",
            f"test:{prefix}:product:labels:lock",
        )
        await client.aclose()
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None or REDIS_URL is None,
    reason="Integration PostgreSQL and Redis URLs are not configured",
)
async def test_add_product_creates_missing_label_and_refreshes_both_caches() -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    service, products, labels, product_cache, label_cache = (
        create_integration_product_service(connection, client, prefix)
    )
    try:
        await create_temp_product_tables(connection)
        result = await service.add(
            AuthorizationObject(uid="integration-user"),
            ProductCreateInfo(
                name="Created with label",
                label_names=" New label ",
                cost=Decimal("10.00"),
                price=Decimal("20.00"),
                stock=3,
            ),
        )

        stored_labels = await labels.list_all()
        stored_products = await products.list_all()
        assert len(stored_labels) == 1
        assert stored_labels[0].name == "New label"
        assert stored_labels[0].created_uid == "integration-user"
        assert len(stored_products) == 1
        assert stored_products[0].label_ids == str(stored_labels[0].id)
        assert await label_cache.get() == {
            "new label": (stored_labels[0].id, "New label")
        }
        assert await product_cache.get_products() == stored_products
        assert result.info.label_names == "New label"
    finally:
        await client.delete(
            f"test:{prefix}:products:info",
            f"test:{prefix}:products:info:version",
            f"test:{prefix}:products:info:lock",
            f"test:{prefix}:product:labels",
            f"test:{prefix}:product:labels:version",
            f"test:{prefix}:product:labels:lock",
        )
        await client.aclose()
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None or REDIS_URL is None,
    reason="Integration PostgreSQL and Redis URLs are not configured",
)
async def test_add_product_failure_rolls_back_new_label(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    service, products, labels, product_cache, label_cache = (
        create_integration_product_service(connection, client, prefix)
    )

    async def fail_insert(*args: object, **kwargs: object) -> object:
        raise ProductRepositoryError("injected product insert failure")

    monkeypatch.setattr(products, "insert", fail_insert)
    try:
        await create_temp_product_tables(connection)
        with pytest.raises(ServiceUnavailableError):
            await service.add(
                AuthorizationObject(uid="integration-user"),
                ProductCreateInfo(
                    name="Must roll back",
                    label_names="New label",
                    cost=Decimal("10.00"),
                    price=Decimal("20.00"),
                    stock=3,
                ),
            )

        assert await labels.list_all() == []
        assert await products.list_all() == []
    finally:
        await client.delete(
            f"test:{prefix}:products:info",
            f"test:{prefix}:products:info:version",
            f"test:{prefix}:products:info:lock",
            f"test:{prefix}:product:labels",
            f"test:{prefix}:product:labels:version",
            f"test:{prefix}:product:labels:lock",
        )
        await client.aclose()
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None or REDIS_URL is None,
    reason="Integration PostgreSQL and Redis URLs are not configured",
)
async def test_add_product_label_cache_failure_keeps_committed_rows_and_invalidates_caches(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    service, products, labels, product_cache, label_cache = (
        create_integration_product_service(connection, client, prefix)
    )

    async def fail_replace(*args: object, **kwargs: object) -> None:
        raise LabelCacheRepositoryError("injected label cache write failure")

    monkeypatch.setattr(label_cache, "replace", fail_replace)
    try:
        await create_temp_product_tables(connection)
        with pytest.raises(ProductCacheUnavailableError):
            await service.add(
                AuthorizationObject(uid="integration-user"),
                ProductCreateInfo(
                    name="Committed product",
                    label_names="New label",
                    cost=Decimal("10.00"),
                    price=Decimal("20.00"),
                    stock=3,
                ),
            )

        assert len(await products.list_all()) == 1
        assert len(await labels.list_all()) == 1
        assert await product_cache.get_products() is None
        assert await label_cache.get() is None
    finally:
        await client.delete(
            f"test:{prefix}:products:info",
            f"test:{prefix}:products:info:version",
            f"test:{prefix}:products:info:lock",
            f"test:{prefix}:product:labels",
            f"test:{prefix}:product:labels:version",
            f"test:{prefix}:product:labels:lock",
        )
        await client.aclose()
        await connection.close()


@pytest.mark.asyncio
@pytest.mark.skipif(
    DATABASE_URL is None or REDIS_URL is None,
    reason="Integration PostgreSQL and Redis URLs are not configured",
)
async def test_legacy_product_snapshot_reloads_stock_from_database() -> None:
    connection = await AsyncConnection.connect(DATABASE_URL)
    client = Redis.from_url(REDIS_URL, decode_responses=True)
    prefix = uuid4().hex
    service, products, labels, product_cache, label_cache = (
        create_integration_product_service(connection, client, prefix)
    )
    product_key = f"test:{prefix}:products:info"
    try:
        await create_temp_product_tables(connection)
        now = datetime(2026, 1, 1)
        async with products.transaction() as transaction:
            label = await labels.insert(
                Label(0, "Stock label", "seed", now, "seed", now),
                transaction,
            )
            product = Product(
                id="1790705105001",
                name="Stock from database",
                label_ids=str(label.id),
                cost=Decimal("10.00"),
                price=Decimal("20.00"),
                stock=23,
                delete_flag=False,
                created_uid="seed",
                created_at=now,
                updated_uid="seed",
                updated_at=now,
            )
            await products.insert(product, transaction)
        async with label_cache.lock() as label_lease:
            await label_cache.replace([label], label_lease)
        async with product_cache.products_lock() as product_lease:
            await product_cache.replace_products(
                [product], {label.id: label.name}, product_lease
            )

        snapshot = json.loads(await client.get(product_key))
        snapshot[0].pop("stock")
        await client.set(product_key, json.dumps(snapshot))

        result = await service.get(
            AuthorizationObject(uid="integration-user"),
            product.id,
        )

        assert result.info.root[0].stock == 23
        refreshed = await product_cache.get_products()
        assert refreshed is not None and refreshed[0].stock == 23
    finally:
        await client.delete(
            product_key,
            f"{product_key}:version",
            f"{product_key}:lock",
            f"test:{prefix}:product:labels",
            f"test:{prefix}:product:labels:version",
            f"test:{prefix}:product:labels:lock",
        )
        await client.aclose()
        await connection.close()
